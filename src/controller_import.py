"""Import controller of Hoopster: the Import page builder, the pure file and paste parsers, the
template builders and the load / publish / discard callbacks (REQ-027, REQ-050; REQ-005
confirmation; task 2026-10-23_drill_import).

Python parses structure only (header, columns, row numbers, size limits); every content rule is
judged by the item type's h_staging check function, so files and pastes get the same messages
(plan D3). A load is one transaction (abandon the page's previous open batch, create, stage,
check); an Import is one transaction (publish, then SET CONSTRAINTS ALL IMMEDIATE). The batch id
comes only from the import-batch store written by the load callbacks; the database re-checks
status on every use (Context Adaptation Table). Every write goes through
app_controller.run_write / run_write_with_reason; every callback is a user-facing boundary
(COD-4-04). Logs carry batch ids, counts and source kinds only, never cell contents.
"""

import base64
import binascii
import csv
import io
import logging
import zlib
from collections.abc import Callable, Iterable, Mapping
from datetime import date, datetime, time
from itertools import repeat
from typing import TypeVar
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from dash import ClientsideFunction, Dash, Input, Output, State, ctx, dcc, no_update
from dash.development.base_component import Component
from dash.exceptions import PreventUpdate
from flask_login import current_user
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.orm import Session

from config.app_config import (
    IMPORT_EXTRA_COLUMNS,
    IMPORT_MAX_BYTES,
    IMPORT_MAX_ROWS,
    IMPORT_MAX_SCANNED_RECORDS,
    IMPORT_XLSX_MAX_MEMBERS,
    IMPORT_XLSX_MAX_UNPACKED_BYTES,
)
from src import data_handler
from src.app_controller import (
    DeviceClock,
    log_failure,
    new_reference_code,
    run_write,
    run_write_with_reason,
)
from src.data_drills import list_reference_options
from src.data_import import (
    IMPORT_ITEM_TYPES,
    NOT_IMPORTABLE_RESULT_TYPES,
    OPEN_STATUSES,
    SOURCE_CSV,
    SOURCE_PASTE,
    SOURCE_XLSX,
    STATUS_READY,
    ImportItemType,
    StagedRow,
    abandon_batch,
    check_batch,
    create_batch,
    get_batch,
    list_batch_rows,
    publish_batch,
    stage_rows,
)
from src.layouts_import import (
    HIDDEN_STYLE,
    SHOWN_STYLE,
    import_layout,
    paste_field_styles,
    preview_items,
    publish_question,
    summary_text,
)

logger = logging.getLogger(__name__)

MESSAGE_FILE_TYPE = "Choose a .csv or .xlsx file."
MESSAGE_FILE_SIZE = "The file is larger than 1 MB."
MESSAGE_FILE_UNREADABLE = "The file could not be read."
MESSAGE_XLSX_UNREADABLE = "The Excel file could not be read."
MESSAGE_UNKNOWN_COLUMNS = "Unknown column(s): {unknown}. The columns are: {columns}."
MESSAGE_DUPLICATE_COLUMN = "Column {column} appears twice."
MESSAGE_MISSING_COLUMN = "The file needs a {column} column."
MESSAGE_NO_ROWS = "The file has no {plural}."
MESSAGE_TOO_MANY_ROWS = (
    f"The file has more than {IMPORT_MAX_ROWS} rows; split it into smaller files."
)
MESSAGE_TOO_MANY_EMPTY_ROWS = "The file has too many empty rows; delete them and load it again."
MESSAGE_PASTE_EMPTY = "Paste at least one drill name."
MESSAGE_PASTE_TOO_MANY = (
    f"The list has more than {IMPORT_MAX_ROWS} lines; split it into smaller lists."
)
MESSAGE_PASTE_CHOICES = "Choose a drill type, difficulty and result type."
MESSAGE_CHOOSE_ITEM = "Choose what to import."
MESSAGE_RELOAD = "Load the file again."
MESSAGE_LOAD_FAILED = "Could not load the import. Reference: {ref}"
MESSAGE_IMPORT_FAILED = "Could not import. Reference: {ref}"
MESSAGE_DISCARD_FAILED = "Could not discard. Reference: {ref}"
MESSAGE_TEMPLATE_FAILED = "Could not build the template. Reference: {ref}"
MESSAGE_DISCARDED = "Discarded."
MESSAGE_IMPORTED = "Imported {count}."

EXTENSION_CSV = ".csv"
EXTENSION_XLSX = ".xlsx"
BASE64_MARKER = ";base64"
CSV_TEXT_ENCODING = "utf-8-sig"
CSV_FALLBACK_ENCODING = "cp1252"
CELL_TRUE = "true"
CELL_FALSE = "false"
PASTE_SEPARATOR = "|"
LIST_SEPARATOR = "; "
TEMPLATE_CSV_NAME = "hoopster-{item}-import-template.csv"
TEMPLATE_XLSX_NAME = "hoopster-{item}-import-template.xlsx"
HELP_SHEET_TITLE = "Help"
HELP_SHEET_HEADER: tuple[str, str] = ("Column", "Rule")
XLSX_STRING_TYPE = "s"
CAPTURE_SHOTS = "shots"
FREE_THROW_VALUE = "FT"
FREE_THROW_LABEL = "Free-throw line"
SPOT_LABEL_SEPARATOR = " · "
IMPORT_BUTTON_LABEL = "Import"
TEMPLATE_CSV_BUTTON = "import-template-csv"
TEMPLATE_XLSX_BUTTON = "import-template-xlsx"

# What a damaged or hostile .xlsx can raise inside zipfile and openpyxl (zip, XML, value
# parsing); each is turned into one message for the user (plan item 16). Sizes are bounded
# before openpyxl runs (_xlsx_archive_ok), so the list is narrowed to file faults (review_1
# Minor): TypeError stays because openpyxl's descriptors raise it for an XML attribute of the
# wrong type; ParseError is the standard-library XML error openpyxl uses (no lxml installed).
XLSX_READ_ERRORS: tuple[type[BaseException], ...] = (
    BadZipFile,
    InvalidFileException,
    LookupError,
    ValueError,
    TypeError,
    OSError,
    EOFError,
    ParseError,
    zlib.error,
)
# NUL cannot be stored in PostgreSQL text, so a CSV holding one is refused at parse time.
CSV_NUL_BYTE = b"\x00"
# Columns read from an Excel sheet, whatever its <dimension> claims (review_1 High).
XLSX_MAX_COLUMNS = (
    max(len(item.columns) for item in IMPORT_ITEM_TYPES.values()) + IMPORT_EXTRA_COLUMNS
)
# One record as a reader yields it: a list of texts (CSV) or a tuple of cell values (Excel).
RawRecord = TypeVar("RawRecord")


# --- Page builder ---

def _paste_options(options: Mapping[str, list[dict]]) -> dict[str, object]:
    """
    Description:
        The paste tab's option lists and the result facts its sections need.

    Parameters:
        options: data_drills.list_reference_options.

    Return Values:
        dict[str, object]: "drill_types" (values = live type NAMES), "difficulties" (codes),
        "result_types" (importable codes only), "spots" (values = spot NUMBERS as text in
        number order, then "FT"), each a list of {"value", "label"}; "result_meta" (code ->
        {"capture_kind", "needs_fixed_seconds"}) of the importable result types.

    Notes:
        Pure. Spot values are spot numbers, never court_spot_id: the staging format the check
        function maps (plan D13).
    """
    importable = [
        row
        for row in options["result_types"]
        if row["result_type_code"] not in NOT_IMPORTABLE_RESULT_TYPES
    ]
    spots = sorted(options["court_spots"], key=lambda spot: int(spot["spot_number"]))
    spot_options = [
        {
            "value": str(int(spot["spot_number"])),
            "label": f"{int(spot['spot_number'])}{SPOT_LABEL_SEPARATOR}{spot['spot_name']}",
        }
        for spot in spots
    ]
    spot_options.append({"value": FREE_THROW_VALUE, "label": FREE_THROW_LABEL})
    return {
        "drill_types": [
            {"value": str(row["type_name"]), "label": str(row["type_name"])}
            for row in options["drill_types"]
        ],
        "difficulties": [
            {"value": str(row["difficulty_code"]), "label": str(row["label"])}
            for row in options["difficulties"]
        ],
        "result_types": [
            {"value": str(row["result_type_code"]), "label": str(row["label"])}
            for row in importable
        ],
        "spots": spot_options,
        "result_meta": {
            str(row["result_type_code"]): {
                "capture_kind": row["capture_kind"],
                "needs_fixed_seconds": bool(row["needs_fixed_seconds"]),
            }
            for row in importable
        },
    }


def build_import_page(
    player_id: int | None, clock: DeviceClock, params: dict[str, str]
) -> Component:
    """
    Description:
        Builds the Import page (REQ-050): item-type select, upload and paste tabs, help and
        the empty result area.

    Parameters:
        player_id: the selected player (unused; the import is not player-scoped).
        clock: the device clock (unused).
        params: the route parameters (none).

    Return Values:
        Component: the page.

    Notes:
        Exceptions propagate to the router, which shows the error page with a reference.
    """
    options = list_reference_options(data_handler.get_session())
    item_types = [
        {"value": code, "label": item.label} for code, item in IMPORT_ITEM_TYPES.items()
    ]
    return import_layout(item_types, _paste_options(options), help_rows(options))


# --- Pure helpers: files (REQ-050, D4, D5, D8) ---

def decode_upload(contents: object, filename: object) -> tuple[str, bytes] | str:
    """
    Description:
        Reads a dcc.Upload value: the source kind from the file name's extension and the
        decoded bytes.

    Parameters:
        contents: import-upload.contents, "data:<type>;base64,<data>".
        filename: import-upload.filename.

    Return Values:
        tuple[str, bytes] | str: ("csv" or "xlsx", the bytes), or a message: "Choose a .csv
        or .xlsx file." / "The file is larger than 1 MB." / "The file could not be read."
        (not base64).

    Notes:
        Pure. The extension is matched case-insensitively. The size is checked on the encoded
        length first, so an oversized upload is never decoded, then on the bytes.
    """
    if not isinstance(filename, str):
        return MESSAGE_FILE_TYPE
    lowered = filename.strip().lower()
    if lowered.endswith(EXTENSION_CSV):
        source_kind = SOURCE_CSV
    elif lowered.endswith(EXTENSION_XLSX):
        source_kind = SOURCE_XLSX
    else:
        return MESSAGE_FILE_TYPE
    if not isinstance(contents, str) or "," not in contents:
        return MESSAGE_FILE_UNREADABLE
    header, encoded = contents.split(",", 1)
    if BASE64_MARKER not in header:
        return MESSAGE_FILE_UNREADABLE
    # Four base64 characters carry three bytes.
    if len(encoded) > (IMPORT_MAX_BYTES // 3 + 1) * 4:
        return MESSAGE_FILE_SIZE
    try:
        data = base64.b64decode(encoded, validate=True)
    except (binascii.Error, ValueError):
        return MESSAGE_FILE_UNREADABLE
    if len(data) > IMPORT_MAX_BYTES:
        return MESSAGE_FILE_SIZE
    return source_kind, data


def _collect_records(
    raw_records: Iterable[RawRecord],
    is_blank: Callable[[RawRecord], bool],
    to_cells: Callable[[RawRecord], list[str]],
    blank_record: list[str],
    source_label: str,
) -> list[list[str]] | str:
    """
    Description:
        Reads a file's records under the import limits (DEF-058). The first record is the
        header and is always kept. A blank record is counted but not kept: a run of blank
        records is stored only when a non-blank record follows it, and then as references to
        the one blank_record, so every kept record stays at its own row number. Trailing
        blank records are never stored.

    Parameters:
        raw_records: the reader's records, in file order.
        is_blank: whether a raw record is blank (every cell empty or white space).
        to_cells: turns a kept raw record into cell texts.
        blank_record: the record that stands for a blank row before a later data row; it is
            shared and never modified.
        source_label: "CSV" or "Excel", for the log line.

    Return Values:
        list[list[str]] | str: the records (header first), or "The file has more than 500
        rows; split it into smaller files." (more than IMPORT_MAX_ROWS non-blank data rows) /
        "The file has too many empty rows; delete them and load it again." (more than
        IMPORT_MAX_SCANNED_RECORDS records read, header and blanks included; one error log,
        counts only).

    Notes:
        Stops at the first limit reached, so neither a sheet of filler rows nor a megabyte of
        blank lines is read to the end (review_1 High). A reader error propagates to the
        caller.
    """
    records: list[list[str]] = []
    scanned = 0
    data_rows = 0
    pending_blanks = 0
    for raw in raw_records:
        scanned += 1
        if scanned > IMPORT_MAX_SCANNED_RECORDS:
            logger.error(
                "%s import file refused: more than %s records read, %s data rows",
                source_label,
                IMPORT_MAX_SCANNED_RECORDS,
                data_rows,
            )
            return MESSAGE_TOO_MANY_EMPTY_ROWS
        if scanned == 1:
            records.append(to_cells(raw))
        elif is_blank(raw):
            pending_blanks += 1
        else:
            data_rows += 1
            if data_rows > IMPORT_MAX_ROWS:
                return MESSAGE_TOO_MANY_ROWS
            records.extend(repeat(blank_record, pending_blanks))
            pending_blanks = 0
            records.append(to_cells(raw))
    return records


def parse_csv_records(data: bytes) -> list[list[str]] | str:
    """
    Description:
        Splits a CSV file into records of cell texts.

    Parameters:
        data: the file's bytes.

    Return Values:
        list[list[str]] | str: the records (header first; a blank line before a later data
        line is kept as [], trailing blank lines are dropped), or "The file could not be
        read." (undecodable, malformed, or holding a NUL byte) / "The file has more than 500
        rows; split it into smaller files." (more than IMPORT_MAX_ROWS non-blank data lines)
        / "The file has too many empty rows; delete them and load it again." (more than
        IMPORT_MAX_SCANNED_RECORDS records read).

    Notes:
        Pure apart from one error log per refusal for an unreadable file or too many records
        (no file contents). UTF-8 (a byte-order mark is dropped), else Windows-1252.
        Comma-separated only (list cells use ";", D5); standard quoting, so a quoted cell
        keeps its line breaks. A NUL byte is refused here because PostgreSQL text cannot hold
        it (review_1 Minor). Limits as in _collect_records (DEF-058).
    """
    if CSV_NUL_BYTE in data:
        logger.error("CSV import file not readable: it contains a NUL byte")
        return MESSAGE_FILE_UNREADABLE
    try:
        text = data.decode(CSV_TEXT_ENCODING)
    except UnicodeDecodeError:
        try:
            text = data.decode(CSV_FALLBACK_ENCODING)
        except UnicodeDecodeError as exc:
            logger.error("CSV import file not readable: %s", type(exc).__name__)
            return MESSAGE_FILE_UNREADABLE
    try:
        return _collect_records(
            csv.reader(io.StringIO(text, newline="")), _is_blank, list, [], "CSV"
        )
    except csv.Error as exc:
        logger.error("CSV import file not readable: %s", type(exc).__name__)
        return MESSAGE_FILE_UNREADABLE


def _cell_text(value: object) -> str:
    """
    Description:
        One Excel cell as text.

    Parameters:
        value: the cell value openpyxl returns.

    Return Values:
        str: "" for None, "true" / "false" for a boolean, a whole float as an integer ("20",
        not "20.0"), dates and times in ISO 8601, anything else as str().

    Notes:
        Pure. bool is tested before numbers because it is an int subclass.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return CELL_TRUE if value else CELL_FALSE
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return str(value)


def _is_blank(cells: list[str]) -> bool:
    """
    Description:
        Whether every cell of a record is empty or white space.

    Parameters:
        cells: the cell texts.

    Return Values:
        bool: True for a blank record.

    Notes:
        Pure.
    """
    return all(not cell.strip() for cell in cells)


def _xlsx_values_blank(values: tuple[object, ...]) -> bool:
    """
    Description:
        Whether every cell value of an Excel row would be an empty or white-space text.

    Parameters:
        values: the row's cell values as openpyxl yields them.

    Return Values:
        bool: True for a blank row (the same verdict as _is_blank on the row's texts).

    Notes:
        Pure. Builds no list, and converts only cells that hold a value, so the many empty
        rows of a padded sheet cost almost nothing (DEF-058).
    """
    return all(value is None or not _cell_text(value).strip() for value in values)


def _xlsx_cells(values: tuple[object, ...]) -> list[str]:
    """
    Description:
        An Excel row as cell texts.

    Parameters:
        values: the row's cell values as openpyxl yields them.

    Return Values:
        list[str]: one text per cell (_cell_text).

    Notes:
        Pure.
    """
    return [_cell_text(value) for value in values]


def _xlsx_archive_ok(data: bytes) -> bool:
    """
    Description:
        Checks an .xlsx file's zip directory before openpyxl opens it: the number of members
        and the total size they would unpack to (review_1 High, zip bomb).

    Parameters:
        data: the file's bytes.

    Return Values:
        bool: True when the archive is readable, has at most IMPORT_XLSX_MAX_MEMBERS members
        and unpacks to at most IMPORT_XLSX_MAX_UNPACKED_BYTES; False otherwise (logged with
        logger.error, counts and the exception class only).

    Notes:
        In memory; nothing is decompressed. The declared sizes are a real bound: the
        standard-library zipfile stops reading a member at its declared size, so a member
        cannot unpack to more than it claims.
    """
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
    except XLSX_READ_ERRORS as exc:
        logger.error("Excel import file not readable: %s", type(exc).__name__)
        return False
    unpacked = sum(member.file_size for member in members)
    if len(members) > IMPORT_XLSX_MAX_MEMBERS or unpacked > IMPORT_XLSX_MAX_UNPACKED_BYTES:
        logger.error(
            "Excel import file refused: %s members, %s bytes unpacked", len(members), unpacked
        )
        return False
    return True


def parse_xlsx_records(data: bytes) -> list[list[str]] | str:
    """
    Description:
        Reads the first worksheet of an .xlsx file into records of cell texts.

    Parameters:
        data: the file's bytes.

    Return Values:
        list[list[str]] | str: the records (header first; every kept record is
        XLSX_MAX_COLUMNS cells wide; a blank row before a later data row is kept as all-empty
        cells, trailing blank rows are dropped), or "The Excel file could not be read."
        (damaged, or over the zip limits of _xlsx_archive_ok) / "The file has more than 500
        rows; split it into smaller files." (more than IMPORT_MAX_ROWS non-blank data rows) /
        "The file has too many empty rows; delete them and load it again." (more than
        IMPORT_MAX_SCANNED_RECORDS rows read).

    Notes:
        Pure (in memory) apart from error logs. Read-only, values only (a formula gives its
        saved result). The sheet's own <dimension> is ignored: at most XLSX_MAX_COLUMNS
        columns are read, and rows are counted as they come, blank or not, so a sheet that
        claims a million rows or columns cannot make the reader pad them (review_1 High).
        Limits as in _collect_records (DEF-058); blank rows are tested on the raw values, so
        no cell texts are built for them. The workbook is always closed. Errors are logged by
        class only.
    """
    if not _xlsx_archive_ok(data):
        return MESSAGE_XLSX_UNREADABLE
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except XLSX_READ_ERRORS as exc:
        logger.error("Excel import file not readable: %s", type(exc).__name__)
        return MESSAGE_XLSX_UNREADABLE
    try:
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()
        return _collect_records(
            sheet.iter_rows(max_col=XLSX_MAX_COLUMNS, values_only=True),
            _xlsx_values_blank,
            _xlsx_cells,
            [""] * XLSX_MAX_COLUMNS,
            "Excel",
        )
    except XLSX_READ_ERRORS as exc:
        logger.error("Excel import file not readable: %s", type(exc).__name__)
        return MESSAGE_XLSX_UNREADABLE
    finally:
        workbook.close()


def _column_key(cell: object) -> str:
    """
    Description:
        A header cell as a column name: trimmed, lower case, spaces and hyphens as
        underscores.

    Parameters:
        cell: the header cell text.

    Return Values:
        str: the column name, e.g. "Drill Name" -> "drill_name".

    Notes:
        Pure.
    """
    return str(cell).strip().lower().replace(" ", "_").replace("-", "_")


def map_columns(header: list[str], item: ImportItemType) -> dict[str, int] | str:
    """
    Description:
        Maps the header row to the item type's staging columns (D5).

    Parameters:
        header: the first record's cells.
        item: the item type.

    Return Values:
        dict[str, int] | str: staging column -> cell index, or a message: "Unknown
        column(s): {list}. The columns are: {columns}." / "Column {x} appears twice." / "The
        file needs a {required_column} column.".

    Notes:
        Pure. Case, surrounding spaces, inner spaces and hyphens are forgiven; trailing empty
        header cells are ignored; a blank header cell between named ones is an unknown
        column. Missing optional columns are staged empty.
    """
    names = [_column_key(cell) for cell in header]
    while names and not names[-1]:
        names.pop()
    known = set(item.columns)
    unknown = [name or "(blank)" for name in names if name not in known]
    if unknown:
        return MESSAGE_UNKNOWN_COLUMNS.format(
            unknown=", ".join(unknown), columns=", ".join(item.columns)
        )
    column_map: dict[str, int] = {}
    for index, name in enumerate(names):
        if name in column_map:
            return MESSAGE_DUPLICATE_COLUMN.format(column=name)
        column_map[name] = index
    if item.required_column not in column_map:
        return MESSAGE_MISSING_COLUMN.format(column=item.required_column)
    return column_map


def file_staged_rows(
    records: list[list[str]], column_map: dict[str, int], item: ImportItemType
) -> list[StagedRow] | str:
    """
    Description:
        Turns the data records of a file into staged rows.

    Parameters:
        records: every record, header first.
        column_map: map_columns' result.
        item: the item type.

    Return Values:
        list[StagedRow] | str: the rows, or "The file has no {plural}." / "The file has more
        than 500 rows; split it into smaller files.".

    Notes:
        Pure. row_no is the spreadsheet row (header = 1, first data row = 2, D8); a record
        whose mapped cells are all blank is skipped but keeps its number. Empty cells and
        missing columns are staged as None; the check function trims the rest.
    """
    rows: list[StagedRow] = []
    for row_no, record in enumerate(records[1:], start=2):
        cells: dict[str, str | None] = {column: None for column in item.columns}
        blank = True
        for column, index in column_map.items():
            text = record[index] if index < len(record) else ""
            if text.strip():
                blank = False
                cells[column] = text
        if blank:
            continue
        rows.append({"row_no": row_no, "cells": cells})
        if len(rows) > IMPORT_MAX_ROWS:
            return MESSAGE_TOO_MANY_ROWS
    if not rows:
        return MESSAGE_NO_ROWS.format(plural=item.plural)
    return rows


def _file_rows(source_kind: str, data: bytes, item: ImportItemType) -> list[StagedRow] | str:
    """
    Description:
        Parses a decoded file into staged rows: records, header mapping, rows.

    Parameters:
        source_kind: "csv" or "xlsx".
        data: the file's bytes.
        item: the item type.

    Return Values:
        list[StagedRow] | str: the rows, or the first file-level message (D4).

    Notes:
        Pure. An empty file has no header, so it gets the missing-column message.
    """
    if source_kind == SOURCE_XLSX:
        records = parse_xlsx_records(data)
    else:
        records = parse_csv_records(data)
    if isinstance(records, str):
        return records
    column_map = map_columns(records[0] if records else [], item)
    if isinstance(column_map, str):
        return column_map
    return file_staged_rows(records, column_map, item)


# --- Pure helpers: paste (REQ-027, D9) ---

def parse_paste_lines(text: object) -> list[tuple[int, str, str | None]] | str:
    """
    Description:
        Splits the pasted list into (line number, name, description), one drill per line:
        "Name" or "Name | description".

    Parameters:
        text: import-paste-text.value.

    Return Values:
        list[tuple[int, str, str | None]] | str: the lines, or "Paste at least one drill
        name." / "The list has more than 500 lines; split it into smaller lists.".

    Notes:
        Pure. Blank lines are skipped but keep their numbers (D8). Only the first "|" splits,
        so a description may contain "|". Name and description are stripped; an empty
        description is None. An empty name is staged as None and reported by the check.
    """
    if not isinstance(text, str):
        return MESSAGE_PASTE_EMPTY
    lines: list[tuple[int, str, str | None]] = []
    for line_no, line in enumerate(text.splitlines(), start=1):
        if not line.strip():
            continue
        name, separator, description = line.partition(PASTE_SEPARATOR)
        lines.append((line_no, name.strip(), description.strip() if separator else None))
        if len(lines) > IMPORT_MAX_ROWS:
            return MESSAGE_PASTE_TOO_MANY
    if not lines:
        return MESSAGE_PASTE_EMPTY
    return [(line_no, name, description or None) for line_no, name, description in lines]


def _seconds_text(seconds: object) -> str | None:
    """
    Description:
        The fixed time as staging text.

    Parameters:
        seconds: import-paste-seconds.value (int, float, text or None).

    Return Values:
        str | None: whole numbers as integer text ("30"), other text as typed, None when
        empty (the check then reports the missing time on every row).

    Notes:
        Pure. A bool is not a number.
    """
    if seconds is None or isinstance(seconds, bool):
        return None
    if isinstance(seconds, int):
        return str(seconds)
    if isinstance(seconds, float):
        return str(int(seconds)) if seconds.is_integer() else str(seconds)
    if isinstance(seconds, str) and seconds.strip():
        return seconds.strip()
    return None


def paste_staged_rows(
    lines: list[tuple[int, str, str | None]],
    primary_type: object,
    difficulty: object,
    result_type: object,
    spots: object,
    seconds: object,
    result_meta: object,
) -> list[StagedRow] | str:
    """
    Description:
        Turns pasted lines into staged drill rows, copying the choices made once onto every
        line (REQ-027).

    Parameters:
        lines: parse_paste_lines' result.
        primary_type: import-paste-type.value (a type NAME).
        difficulty: import-paste-difficulty.value (a difficulty code).
        result_type: import-paste-result.value (a result type code).
        spots: import-paste-spots.value: spot NUMBERS as text and "FT", in picked order.
        seconds: import-paste-seconds.value.
        result_meta: the import-result-meta store (code -> {"capture_kind",
            "needs_fixed_seconds"}).

    Return Values:
        list[StagedRow] | str: the rows, or "Choose a drill type, difficulty and result
        type." when any of the three is empty.

    Notes:
        Pure. Spots ("16; FT") only for a shooting result type and fixed_seconds only for a
        type that needs it (empty there stays None, so the check reports it per row); every
        other column is None.
    """
    choices = (primary_type, difficulty, result_type)
    if not all(isinstance(choice, str) and choice.strip() for choice in choices):
        return MESSAGE_PASTE_CHOICES
    meta = result_meta if isinstance(result_meta, dict) else {}
    facts = meta.get(result_type)
    facts = facts if isinstance(facts, dict) else {}
    spots_text: str | None = None
    if facts.get("capture_kind") == CAPTURE_SHOTS and isinstance(spots, list):
        picked = [str(spot) for spot in spots if isinstance(spot, str) and spot]
        spots_text = LIST_SEPARATOR.join(picked) or None
    seconds_value = _seconds_text(seconds) if facts.get("needs_fixed_seconds") else None
    return [
        {
            "row_no": line_no,
            "cells": {
                "drill_name": name or None,
                "description": description,
                "primary_type": str(primary_type),
                "difficulty": str(difficulty),
                "result_type": str(result_type),
                "spots": spots_text,
                "fixed_seconds": seconds_value,
            },
        }
        for line_no, name, description in lines
    ]


# --- Pure helpers: help and templates (D5) ---

def help_rows(options: Mapping[str, list[dict]]) -> list[tuple[str, str]]:
    """
    Description:
        The column help of the drill import: one (column, rule) per staging column, in
        template order.

    Parameters:
        options: data_drills.list_reference_options.

    Return Values:
        list[tuple[str, str]]: e.g. ("drill_name", "Required. 1–100 characters; must not
        match an existing drill.").

    Notes:
        Pure. Lists the live values, so the help never names a type or label that is gone;
        the rules mirror h_staging.fn_check_drill_batch.
    """
    importable = [
        row
        for row in options["result_types"]
        if row["result_type_code"] not in NOT_IMPORTABLE_RESULT_TYPES
    ]
    fixed_labels = [str(row["label"]) for row in importable if row["needs_fixed_seconds"]]
    spot_numbers = [int(spot["spot_number"]) for spot in options["court_spots"]]
    spot_range = (
        f"{min(spot_numbers)}–{max(spot_numbers)}" if spot_numbers else "of the court"
    )

    def labels(rows: list[dict], key: str) -> str:
        """
        Description:
            The labels of option rows, comma-separated.

        Parameters:
            rows: the option rows.
            key: the label key.

        Return Values:
            str: "A, B, C".

        Notes:
            Pure.
        """
        return ", ".join(str(row[key]) for row in rows)

    return [
        ("drill_name", "Required. 1–100 characters; must not match an existing drill."),
        ("description", "Optional, up to 1000 characters."),
        ("primary_type", f"Required. One of: {labels(options['drill_types'], 'type_name')}."),
        ("secondary_types", "Optional. Type names separated by ;"),
        ("difficulty", f"Required. One of: {labels(options['difficulties'], 'label')}."),
        ("result_type", f"Required. One of: {labels(importable, 'label')}."),
        (
            "symmetry_type",
            f"Optional. One of: {labels(options['symmetry_types'], 'label')}.",
        ),
        ("default_target", "Optional whole number 1–99999."),
        (
            "fixed_seconds",
            f"Required for {' / '.join(fixed_labels) or 'fixed-time'} drills only; "
            "whole seconds 1–99999.",
        ),
        (
            "spots",
            f"Shooting only. Spot numbers {spot_range} or FT, separated by ;",
        ),
    ]


def template_csv_bytes(item: ImportItemType) -> bytes:
    """
    Description:
        The CSV template: the header row only (example rows would be imported by mistake,
        D5).

    Parameters:
        item: the item type.

    Return Values:
        bytes: UTF-8 with a byte-order mark, the item type's columns comma-separated.

    Notes:
        Pure.
    """
    buffer = io.StringIO()
    csv.writer(buffer).writerow(item.columns)
    return buffer.getvalue().encode(CSV_TEXT_ENCODING)


def _write_text_row(sheet: Worksheet, row_no: int, values: tuple[str, ...]) -> None:
    """
    Description:
        Writes one row of texts, each cell typed as a string.

    Parameters:
        sheet: the worksheet.
        row_no: the 1-based row number.
        values: the texts.

    Return Values:
        None.

    Notes:
        The explicit string type is set after the value, because openpyxl types a text
        starting with "=" as a formula; so such a text is never written as one.
    """
    for column_no, value in enumerate(values, start=1):
        cell = sheet.cell(row=row_no, column=column_no)
        cell.value = value
        cell.data_type = XLSX_STRING_TYPE


def template_xlsx_bytes(item: ImportItemType, help_rows_list: list[tuple[str, str]]) -> bytes:
    """
    Description:
        The Excel template: sheet "Drills" (the item type's label) with the header row, and
        sheet "Help" with (Column, Rule) per column.

    Parameters:
        item: the item type.
        help_rows_list: help_rows' result.

    Return Values:
        bytes: the .xlsx file.

    Notes:
        Built in memory, nothing on disk (COD-20-04). Every cell is written as text.
    """
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = item.label
    _write_text_row(data_sheet, 1, item.columns)
    help_sheet = workbook.create_sheet(HELP_SHEET_TITLE)
    _write_text_row(help_sheet, 1, HELP_SHEET_HEADER)
    for row_no, (column, rule) in enumerate(help_rows_list, start=2):
        _write_text_row(help_sheet, row_no, (column, rule))
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# --- Shared callback helpers ---

def _item_words(item: ImportItemType) -> dict[str, str]:
    """
    Description:
        The item type's words for the layout's pure text helpers.

    Parameters:
        item: the item type.

    Return Values:
        dict[str, str]: {"code", "label", "singular", "plural"}.

    Notes:
        Pure.
    """
    return {
        "code": item.code,
        "label": item.label,
        "singular": item.singular,
        "plural": item.plural,
    }


def _count_words(count: int, item: ImportItemType) -> str:
    """
    Description:
        "1 drill" / "3 drills".

    Parameters:
        count: the number.
        item: the item type.

    Return Values:
        str: the count and the right word.

    Notes:
        Pure.
    """
    return f"{count} {item.singular if count == 1 else item.plural}"


def _item_of(item_type: object) -> ImportItemType | None:
    """
    Description:
        The registry entry of the item-type select's value.

    Parameters:
        item_type: import-item-type.value.

    Return Values:
        ImportItemType | None: the entry, or None for anything not a registry key.

    Notes:
        Pure.
    """
    return IMPORT_ITEM_TYPES.get(item_type) if isinstance(item_type, str) else None


def _stored_batch(batch_store: object) -> tuple[int, ImportItemType, int] | None:
    """
    Description:
        Reads the import-batch store.

    Parameters:
        batch_store: import-batch.data, {"import_batch_id", "item_type", "row_count"}.

    Return Values:
        tuple[int, ImportItemType, int] | None: (BATCH id, item type, row count), or None
        for an empty or malformed store.

    Notes:
        Pure. The client can alter the store; the database re-checks the batch's status on
        every use.
    """
    if not isinstance(batch_store, dict):
        return None
    batch_id = batch_store.get("import_batch_id")
    row_count = batch_store.get("row_count")
    item = _item_of(batch_store.get("item_type"))
    if not isinstance(batch_id, int) or isinstance(batch_id, bool) or item is None:
        return None
    if not isinstance(row_count, int) or isinstance(row_count, bool):
        row_count = 0
    return batch_id, item, row_count


def _preview_labels(options: Mapping[str, list[dict]]) -> dict[str, dict[str, str]]:
    """
    Description:
        The code -> label maps of the preview's detail line.

    Parameters:
        options: data_drills.list_reference_options.

    Return Values:
        dict[str, dict[str, str]]: {"difficulties": {code: label}, "result_types": {code:
        label}}.

    Notes:
        Pure.
    """
    return {
        "difficulties": {
            str(row["difficulty_code"]): str(row["label"]) for row in options["difficulties"]
        },
        "result_types": {
            str(row["result_type_code"]): str(row["label"]) for row in options["result_types"]
        },
    }


def _preview_outputs(
    batch: dict | None, rows: list[dict], labels: dict, feedback: str
) -> tuple[object, ...]:
    """
    Description:
        The result area for a batch, in the callbacks' output order.

    Parameters:
        batch: data_import.get_batch, or None to clear the area.
        rows: data_import.list_batch_rows of the batch.
        labels: _preview_labels.
        feedback: the import-feedback text.

    Return Values:
        tuple[object, ...]: import-batch.data ({"import_batch_id", "item_type",
        "row_count"} while the batch is open, else None), import-summary, import-preview,
        import-confirm.disabled (True unless ready), import-confirm.children ("Import {n}
        drills" when ready, else "Import"), import-discard.disabled (True unless open),
        import-feedback, import-done.style (hidden), import-done-text ("").

    Notes:
        Pure.
    """
    item = _item_of(batch.get("item_type")) if batch else None
    if batch is None or item is None:
        return None, "", [], True, IMPORT_BUTTON_LABEL, True, feedback, HIDDEN_STYLE, ""
    status = batch.get("batch_status")
    row_count = int(batch.get("row_count") or 0)
    is_open = status in OPEN_STATUSES
    is_ready = status == STATUS_READY
    store = (
        {
            "import_batch_id": int(batch["import_batch_id"]),
            "item_type": item.code,
            "row_count": row_count,
        }
        if is_open
        else None
    )
    confirm_label = (
        f"{IMPORT_BUTTON_LABEL} {_count_words(row_count, item)}"
        if is_ready
        else IMPORT_BUTTON_LABEL
    )
    return (
        store,
        summary_text(batch, _item_words(item)),
        preview_items(rows, status, labels),
        not is_ready,
        confirm_label,
        not is_open,
        feedback,
        HIDDEN_STYLE,
        "",
    )


def _message_outputs(feedback: str) -> tuple[object, ...]:
    """
    Description:
        Result-area outputs that change only the feedback line (the preview stays as it is).

    Parameters:
        feedback: the message.

    Return Values:
        tuple[object, ...]: no_update for every output of _preview_outputs except
        import-feedback.

    Notes:
        Pure.
    """
    return (
        no_update, no_update, no_update, no_update, no_update, no_update,
        feedback, no_update, no_update,
    )


def _failure_message(reason: str | None, ref: str | None, template: str) -> str:
    """
    Description:
        The text a failed write shows: the database's reason when there is one, else the
        template with the reference.

    Parameters:
        reason: the message of run_write_with_reason, or None.
        ref: the reference code.
        template: a message with a "{ref}" placeholder.

    Return Values:
        str: the message.

    Notes:
        Pure.
    """
    return reason if reason else template.format(ref=ref)


def _read_preview(
    item: ImportItemType, batch_id: int, feedback: str
) -> tuple[tuple[object, ...], dict | None]:
    """
    Description:
        Reads a batch and its rows and builds the result area.

    Parameters:
        item: the item type of the batch.
        batch_id: the BATCH id.
        feedback: the import-feedback text to show with the preview.

    Return Values:
        tuple[tuple[object, ...], dict | None]: _preview_outputs and the batch read (None
        when it is gone or the read failed). On a failed read the store keeps the batch (so
        a later load still abandons it and Discard works), the rest is cleared, and the
        feedback is "Could not load the import. Reference: {ref}".

    Notes:
        A boundary helper (COD-4-04).
    """
    try:
        session = data_handler.get_session()
        batch = get_batch(session, batch_id)
        rows = list_batch_rows(session, item.code, batch_id)
        labels = _preview_labels(list_reference_options(session))
    except Exception as exc:
        ref = new_reference_code()
        log_failure(
            logger, exc, "Import preview read failed for batch %s (ref %s)", batch_id, ref
        )
        store = {"import_batch_id": batch_id, "item_type": item.code, "row_count": 0}
        cleared = _preview_outputs(None, [], {}, MESSAGE_LOAD_FAILED.format(ref=ref))
        return (store, *cleared[1:5], False, *cleared[6:]), None
    return _preview_outputs(batch, rows, labels, feedback), batch


def _load_batch(
    item: ImportItemType,
    source_kind: str,
    source_name: str | None,
    rows: list[StagedRow],
    batch_store: object,
) -> tuple[object, ...]:
    """
    Description:
        Loads rows as a new batch in one transaction: abandons the page's previous open batch
        (D10), creates the batch, stages every row and runs the check; then shows the
        preview.

    Parameters:
        item: the item type.
        source_kind: "csv", "xlsx" or "paste".
        source_name: the file name, or None for a paste.
        rows: the staged rows (at least one).
        batch_store: import-batch.data (the previous batch, if any).

    Return Values:
        tuple[object, ...]: _preview_outputs; on a failed load only the feedback changes
        (the database's reason, else "Could not import. Reference: {ref}") and the old
        preview stays, because its abandonment was rolled back too.

    Notes:
        A boundary (COD-4-04). The acting user is current_user.user_id.
    """
    previous = _stored_batch(batch_store)

    def work(session: Session) -> int:
        """
        Description:
            The load transaction.

        Parameters:
            session: the write session.

        Return Values:
            int: the new BATCH id.

        Notes:
            Four statements at most (abandon, create, stage, check), none in a loop.
        """
        if previous is not None:
            abandon_batch(session, previous[0])
        batch_id = create_batch(session, item.code, source_kind, source_name)
        stage_rows(session, item.code, batch_id, rows)
        check_batch(session, item.code, batch_id)
        return batch_id

    result, reason = run_write_with_reason(current_user.user_id, work, "load import")
    if not result.ok or result.value is None:
        return _message_outputs(_failure_message(reason, result.ref, MESSAGE_IMPORT_FAILED))
    batch_id = result.value
    outputs, batch = _read_preview(item, batch_id, "")
    if batch is not None:
        logger.info(
            "Loaded import batch %s (%s, %s rows, %s errors)",
            batch_id,
            source_kind,
            batch.get("row_count"),
            batch.get("error_row_count"),
        )
    return outputs


# --- Callbacks: load (REQ-050, REQ-027) ---

def load_file(
    contents: object, filename: object, item_type: object, batch_store: object
) -> tuple[object, ...]:
    """
    Description:
        A file was chosen: checks its type, size and structure, then loads it as a new batch
        and shows the preview.

    Parameters:
        contents: import-upload.contents.
        filename: import-upload.filename.
        item_type: import-item-type.value.
        batch_store: import-batch.data.

    Return Values:
        tuple[object, ...]: _preview_outputs, then import-upload.contents (None, so the same
        file can be chosen again).

    Notes:
        A boundary (COD-4-04). A file-level problem (D4) only sets the feedback: nothing is
        staged and the preview is untouched.
    """
    if contents is None:
        raise PreventUpdate
    item = _item_of(item_type)
    if item is None:
        return (*_message_outputs(MESSAGE_CHOOSE_ITEM), None)
    decoded = decode_upload(contents, filename)
    if isinstance(decoded, str):
        return (*_message_outputs(decoded), None)
    source_kind, data = decoded
    rows = _file_rows(source_kind, data, item)
    if isinstance(rows, str):
        return (*_message_outputs(rows), None)
    source_name = filename.strip() if isinstance(filename, str) else None
    return (*_load_batch(item, source_kind, source_name, rows, batch_store), None)


def load_paste(
    n_clicks: int | None,
    item_type: object,
    text: object,
    primary_type: object,
    difficulty: object,
    result_type: object,
    spots: object,
    seconds: object,
    batch_store: object,
    result_meta: object,
) -> tuple[object, ...]:
    """
    Description:
        Check list: turns the pasted lines into a new batch with the chosen type,
        difficulty, result type (and spots or fixed time) and shows the preview (REQ-027).

    Parameters:
        n_clicks: import-paste-check.n_clicks.
        item_type: import-item-type.value.
        text: import-paste-text.value.
        primary_type: import-paste-type.value.
        difficulty: import-paste-difficulty.value.
        result_type: import-paste-result.value.
        spots: import-paste-spots.value.
        seconds: import-paste-seconds.value.
        batch_store: import-batch.data.
        result_meta: import-result-meta.data.

    Return Values:
        tuple[object, ...]: _preview_outputs.

    Notes:
        A boundary (COD-4-04). Source kind "paste", no source name.
    """
    if not n_clicks:
        raise PreventUpdate
    item = _item_of(item_type)
    if item is None:
        return _message_outputs(MESSAGE_CHOOSE_ITEM)
    lines = parse_paste_lines(text)
    if isinstance(lines, str):
        return _message_outputs(lines)
    rows = paste_staged_rows(
        lines, primary_type, difficulty, result_type, spots, seconds, result_meta
    )
    if isinstance(rows, str):
        return _message_outputs(rows)
    return _load_batch(item, SOURCE_PASTE, None, rows, batch_store)


def show_paste_fields(result_type: object, result_meta: object) -> tuple[dict, dict]:
    """
    Description:
        Shows the spots section for a shooting result type and the fixed-time section for a
        type that needs it.

    Parameters:
        result_type: import-paste-result.value.
        result_meta: import-result-meta.data.

    Return Values:
        tuple[dict, dict]: the styles of import-paste-spots-section and
        import-paste-seconds-section.

    Notes:
        Pure apart from Dash; no database.
    """
    meta = result_meta if isinstance(result_meta, dict) else {}
    return paste_field_styles(result_type if isinstance(result_type, str) else None, meta)


# --- Callbacks: Import and Discard (REQ-050, REQ-005, D6, D11) ---

def open_publish(n_clicks: int | None, batch_store: object) -> tuple[bool, str, str]:
    """
    Description:
        Import: opens the confirmation with the number of items to add.

    Parameters:
        n_clicks: import-confirm.n_clicks.
        batch_store: import-batch.data.

    Return Values:
        tuple[bool, str, str]: import-publish-modal.opened, import-publish-text and
        import-publish-feedback ("").

    Notes:
        No database access.
    """
    if not n_clicks:
        raise PreventUpdate
    stored = _stored_batch(batch_store)
    if stored is None:
        raise PreventUpdate
    _, item, row_count = stored
    return True, publish_question(row_count, _item_words(item)), ""


def _after_failed_publish(
    item: ImportItemType, batch_id: int, feedback: str
) -> tuple[object, ...]:
    """
    Description:
        After a refused or failed publish: re-runs the check on a still-open batch, so the
        preview shows what changed (D11), and re-renders it with the message.

    Parameters:
        item: the item type.
        batch_id: the BATCH id.
        feedback: the database's reason, or the reference message.

    Return Values:
        tuple[object, ...]: _preview_outputs (the failure message as feedback).

    Notes:
        A boundary helper (COD-4-04). The re-check's own failure is only logged (by
        run_write); the preview then shows the batch as it is.
    """
    try:
        batch = get_batch(data_handler.get_session(), batch_id)
    except Exception as exc:
        ref = new_reference_code()
        log_failure(logger, exc, "Import batch read failed for %s (ref %s)", batch_id, ref)
        return _message_outputs(feedback)
    if batch is not None and batch.get("batch_status") in OPEN_STATUSES:
        run_write(
            current_user.user_id,
            lambda session: check_batch(session, item.code, batch_id),
            "re-check import",
        )
    outputs, _ = _read_preview(item, batch_id, feedback)
    return outputs


def confirm_publish(
    n_clicks: int | None, batch_store: object, item_type: object
) -> tuple[object, ...]:
    """
    Description:
        The confirmation's Import: publishes the batch all or nothing.

    Parameters:
        n_clicks: import-publish-confirm.n_clicks.
        batch_store: import-batch.data.
        item_type: import-item-type.value; must equal the store's item type.

    Return Values:
        tuple[object, ...]: import-publish-modal.opened (False), then _preview_outputs:
        cleared with import-done shown ("Imported {n} drills.") on success; on failure the
        re-checked preview with the database's message (e.g. "Batch 7 has 1 rows with
        errors; nothing was published") or "Could not import. Reference: {ref}".

    Notes:
        A boundary (COD-4-04). The acting user is current_user.user_id.
    """
    if not n_clicks:
        raise PreventUpdate
    stored = _stored_batch(batch_store)
    if stored is None or _item_of(item_type) is not stored[1]:
        return (False, *_message_outputs(MESSAGE_RELOAD))
    batch_id, item, _ = stored
    result, reason = run_write_with_reason(
        current_user.user_id,
        lambda session: publish_batch(session, item.code, batch_id),
        "publish import",
    )
    if not result.ok or result.value is None:
        feedback = _failure_message(reason, result.ref, MESSAGE_IMPORT_FAILED)
        return (False, *_after_failed_publish(item, batch_id, feedback))
    published = result.value
    logger.info("Published import batch %s (%s rows)", batch_id, published)
    cleared = _preview_outputs(None, [], {}, "")
    done_text = MESSAGE_IMPORTED.format(count=_count_words(published, item))
    return (False, *cleared[:7], SHOWN_STYLE, done_text)


def discard_import(n_clicks: int | None, batch_store: object) -> tuple[object, ...]:
    """
    Description:
        Discard: marks the previewed batch abandoned and clears the preview (no
        confirmation: nothing is deleted, D6).

    Parameters:
        n_clicks: import-discard.n_clicks.
        batch_store: import-batch.data.

    Return Values:
        tuple[object, ...]: _preview_outputs, cleared, with feedback "Discarded."; on a
        failed write only the feedback changes ("Could not discard. Reference: {ref}").

    Notes:
        A boundary (COD-4-04). A batch that was already published or abandoned is simply
        cleared from the page.
    """
    if not n_clicks:
        raise PreventUpdate
    stored = _stored_batch(batch_store)
    if stored is None:
        raise PreventUpdate
    batch_id = stored[0]
    result = run_write(
        current_user.user_id,
        lambda session: abandon_batch(session, batch_id),
        "discard import",
    )
    if not result.ok:
        return _message_outputs(MESSAGE_DISCARD_FAILED.format(ref=result.ref))
    logger.info("Discarded import batch %s (changed: %s)", batch_id, result.value)
    return _preview_outputs(None, [], {}, MESSAGE_DISCARDED)


def download_template(
    n_csv: int | None, n_xlsx: int | None, item_type: object
) -> tuple[object, object]:
    """
    Description:
        Download CSV template / Download Excel template.

    Parameters:
        n_csv: import-template-csv.n_clicks.
        n_xlsx: import-template-xlsx.n_clicks.
        item_type: import-item-type.value.

    Return Values:
        tuple[object, object]: import-download.data (dcc.send_bytes,
        "hoopster-drill-import-template.csv" / ".xlsx") and import-feedback.

    Notes:
        A boundary (COD-4-04): the Excel template reads the live option lists for its Help
        sheet. Which button by ctx.triggered_id.
    """
    triggered = ctx.triggered_id
    clicks = {TEMPLATE_CSV_BUTTON: n_csv, TEMPLATE_XLSX_BUTTON: n_xlsx}
    if triggered not in clicks or not clicks[triggered]:
        raise PreventUpdate
    item = _item_of(item_type)
    if item is None:
        return no_update, MESSAGE_CHOOSE_ITEM
    if triggered == TEMPLATE_CSV_BUTTON:
        csv_name = TEMPLATE_CSV_NAME.format(item=item.code)
        return dcc.send_bytes(template_csv_bytes(item), csv_name), ""
    try:
        options = list_reference_options(data_handler.get_session())
        content = template_xlsx_bytes(item, help_rows(options))
    except Exception as exc:
        ref = new_reference_code()
        log_failure(logger, exc, "Import template build failed (ref %s)", ref)
        return no_update, MESSAGE_TEMPLATE_FAILED.format(ref=ref)
    return dcc.send_bytes(content, TEMPLATE_XLSX_NAME.format(item=item.code)), ""


# --- Registration ---

def _preview_output_list(duplicate: bool) -> list[Output]:
    """
    Description:
        The result-area Outputs in _preview_outputs order.

    Parameters:
        duplicate: True for every writer after load_file (allow_duplicate).

    Return Values:
        list[Output]: nine Outputs.

    Notes:
        A flag rather than two lists, so the order is written once (COD-13-05 is advisory).
    """
    return [
        Output("import-batch", "data", allow_duplicate=duplicate),
        Output("import-summary", "children", allow_duplicate=duplicate),
        Output("import-preview", "children", allow_duplicate=duplicate),
        Output("import-confirm", "disabled", allow_duplicate=duplicate),
        Output("import-confirm", "children", allow_duplicate=duplicate),
        Output("import-discard", "disabled", allow_duplicate=duplicate),
        Output("import-feedback", "children", allow_duplicate=duplicate),
        Output("import-done", "style", allow_duplicate=duplicate),
        Output("import-done-text", "children", allow_duplicate=duplicate),
    ]


def register_import_callbacks(app: Dash) -> None:
    """
    Description:
        Registers the Import page callbacks: file load, paste load, the paste sections,
        Import (open, cancel, confirm), Discard and the template downloads.

    Parameters:
        app: the Dash application.

    Return Values:
        None.

    Notes:
        load_file is the first writer of the result area; every later writer uses
        allow_duplicate. Every server callback has prevent_initial_call and ignores a run
        without a click (router insertion trap, notes_for_agents §8). Nothing is added to
        PUBLIC_CALLBACK_OUTPUT_IDS.
    """
    app.callback(
        *_preview_output_list(False),
        Output("import-upload", "contents"),
        Input("import-upload", "contents"),
        State("import-upload", "filename"),
        State("import-item-type", "value"),
        State("import-batch", "data"),
        prevent_initial_call=True,
    )(load_file)
    app.callback(
        *_preview_output_list(True),
        Input("import-paste-check", "n_clicks"),
        State("import-item-type", "value"),
        State("import-paste-text", "value"),
        State("import-paste-type", "value"),
        State("import-paste-difficulty", "value"),
        State("import-paste-result", "value"),
        State("import-paste-spots", "value"),
        State("import-paste-seconds", "value"),
        State("import-batch", "data"),
        State("import-result-meta", "data"),
        prevent_initial_call=True,
    )(load_paste)
    app.callback(
        Output("import-paste-spots-section", "style"),
        Output("import-paste-seconds-section", "style"),
        Input("import-paste-result", "value"),
        State("import-result-meta", "data"),
        prevent_initial_call=True,
    )(show_paste_fields)
    app.callback(
        Output("import-publish-modal", "opened"),
        Output("import-publish-text", "children"),
        Output("import-publish-feedback", "children"),
        Input("import-confirm", "n_clicks"),
        State("import-batch", "data"),
        prevent_initial_call=True,
    )(open_publish)
    app.clientside_callback(
        ClientsideFunction("hoopster", "close_modal"),
        Output("import-publish-modal", "opened", allow_duplicate=True),
        Input("import-publish-cancel", "n_clicks"),
        prevent_initial_call=True,
    )
    app.callback(
        Output("import-publish-modal", "opened", allow_duplicate=True),
        *_preview_output_list(True),
        Input("import-publish-confirm", "n_clicks"),
        State("import-batch", "data"),
        State("import-item-type", "value"),
        prevent_initial_call=True,
    )(confirm_publish)
    app.callback(
        *_preview_output_list(True),
        Input("import-discard", "n_clicks"),
        State("import-batch", "data"),
        prevent_initial_call=True,
    )(discard_import)
    app.callback(
        Output("import-download", "data"),
        Output("import-feedback", "children", allow_duplicate=True),
        Input(TEMPLATE_CSV_BUTTON, "n_clicks"),
        Input(TEMPLATE_XLSX_BUTTON, "n_clicks"),
        State("import-item-type", "value"),
        prevent_initial_call=True,
    )(download_template)
